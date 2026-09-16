#!/usr/bin/env python3
"""Los cinco documentos de un pack sectorial.

El formato lo pide el pack: dos Word editables, dos Excel y un PDF. Editables a propósito —el
cliente los adapta a su expediente— salvo el caso resuelto, que se lee y no se rellena.

Todo lo normativo entra por parámetro desde `fuentes.py`, que es donde están las citas con su
artículo. Aquí solo se maqueta: si hiciera falta cambiar una cita, se cambia allí y sale en los
seis packs a la vez, en vez de en treinta ficheros.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib.enums import TA_JUSTIFY
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import ListFlowable, ListItem, PageBreak, Paragraph, SimpleDocTemplate, Spacer

from . import fuentes

AZUL = RGBColor(0x1E, 0x3A, 0x8A)
CABECERA = PatternFill("solid", fgColor="1E3A8A")


# ── portada común ────────────────────────────────────────────────────────────────────────
#
# La misma en los cinco documentos y en los seis packs: versión, fecha de verificación y aviso.
# La fecha es lo que da o quita valor al documento, así que va arriba y no en un pie.


def _portada_docx(doc: Document, titulo: str, sector: str) -> None:
    marca = doc.add_paragraph()
    marca.add_run("INKOOVA ACADEMY").bold = True
    marca.runs[0].font.size = Pt(9)
    marca.runs[0].font.color.rgb = AZUL

    encabezado = doc.add_heading(titulo, level=0)
    encabezado.alignment = WD_ALIGN_PARAGRAPH.LEFT

    ficha = doc.add_paragraph()
    ficha.add_run(f"Pack sectorial · {sector}\n").bold = True
    ficha.add_run(f"Versión {fuentes.VERSION}\n")
    ficha.add_run(f"Verificación normativa: {fuentes.FECHA_VERIFICACION}").bold = True

    aviso = doc.add_paragraph()
    nota = aviso.add_run(fuentes.AVISO)
    nota.italic = True
    nota.font.size = Pt(9)

    doc.add_paragraph()


def _cita(doc: Document, referencia: str, texto: str) -> None:
    """Un requisito con su artículo delante. El artículo va en negrita porque es lo que se busca."""
    parrafo = doc.add_paragraph(style="List Bullet")
    parrafo.add_run(f"{referencia} · ").bold = True
    parrafo.add_run(texto)


# ── 1 · checklist de clasificación ───────────────────────────────────────────────────────


def checklist_clasificacion(destino: Path, pack) -> None:
    """
    El árbol de decisión del artículo 6: ¿es de alto riesgo o no?

    Lo que lo hace útil no es la lista del Anexo III —esa está en cualquier resumen— sino el
    filtro del 6.3 y su excepción. Mucha gente se queda en «está en el Anexo III, luego es de
    alto riesgo» y se salta que puede no serlo; y quien conoce el filtro se salta a menudo que
    la elaboración de perfiles lo desactiva.
    """
    doc = Document()
    _portada_docx(doc, "Checklist de clasificación de riesgo", pack.sector)

    doc.add_paragraph(
        "Se rellena por SISTEMA, no por empresa. Un mismo departamento puede tener un sistema de "
        "alto riesgo y otro que no lo sea, y la clasificación es de cada uno."
    )

    doc.add_heading("Paso 1 · Identificación del sistema", level=1)
    for campo in ("Nombre del sistema", "Finalidad prevista (una frase)",
                  "¿Quién lo desarrolla?", "¿Quién lo despliega?",
                  "¿Sobre qué personas produce efectos?", "Responsable de esta ficha", "Fecha"):
        doc.add_paragraph(f"{campo}: ______________________________________________")

    doc.add_heading("Paso 2 · ¿Es un producto regulado?", level=1)
    doc.add_paragraph(
        f"Según el art. 6.1 del {fuentes.AI_ACT}, un sistema es de alto riesgo cuando es "
        "componente de seguridad de un producto cubierto por la legislación de armonización del "
        "Anexo I —o es él mismo ese producto— Y ese producto debe someterse a evaluación de la "
        "conformidad por un tercero."
    )
    doc.add_paragraph("☐ Sí, las dos condiciones → ALTO RIESGO. No hace falta seguir.")
    doc.add_paragraph("☐ No → continúa en el paso 3.")

    doc.add_heading("Paso 3 · ¿Está en alguno de los ámbitos del Anexo III?", level=1)
    doc.add_paragraph(
        "Marca solo lo que describa la finalidad PREVISTA del sistema. Los ámbitos relevantes "
        "para este sector van primero; el resto se incluye para poder descartar con fundamento."
    )

    doc.add_heading("Ámbitos de este sector", level=2)
    for punto in pack.anexo_iii:
        _cita(doc, f"Anexo III.{punto}", fuentes.ANEXO_III[punto])

    if pack.matiz_anexo:
        aviso = doc.add_paragraph()
        aviso.add_run("Atención: ").bold = True
        aviso.add_run(pack.matiz_anexo)

    doc.add_heading("Resto de ámbitos", level=2)
    for punto, texto in fuentes.ANEXO_III.items():
        if punto not in pack.anexo_iii:
            _cita(doc, f"Anexo III.{punto}", texto)

    doc.add_paragraph("☐ Ninguno → el sistema NO es de alto riesgo por el Anexo III.")

    doc.add_heading("Paso 4 · El filtro del artículo 6.3", level=1)
    doc.add_paragraph(
        "Estar en el Anexo III no basta. Un sistema de esa lista NO se considera de alto riesgo "
        "cuando no plantea un riesgo importante de causar un perjuicio a la salud, la seguridad o "
        "los derechos fundamentales, también por no influir sustancialmente en el resultado de la "
        "toma de decisiones, y se cumple CUALQUIERA de estas condiciones:"
    )
    for letra, texto in fuentes.FILTRO_ART6_3:
        _cita(doc, f"art. 6.3.{letra}", texto)

    corte = doc.add_paragraph()
    corte.add_run("Y aquí está el corte que casi nadie aplica: ").bold = True
    corte.add_run(fuentes.EXCEPCION_PERFILES)

    doc.add_paragraph("☐ Se cumple alguna condición y NO hay elaboración de perfiles → no es de alto riesgo.")
    doc.add_paragraph("☐ No se cumple ninguna, o hay elaboración de perfiles → ALTO RIESGO.")

    doc.add_heading("Paso 5 · Documentar la conclusión", level=1)
    doc.add_paragraph(fuentes.DOCUMENTAR_NO_ALTO_RIESGO)
    doc.add_paragraph(
        "Justificación de la conclusión (por qué se cumple la condición marcada, con referencia "
        "a cómo funciona el sistema y no solo a lo que se pretende de él):"
    )
    for _ in range(6):
        doc.add_paragraph("_____________________________________________________________")

    doc.add_heading("Fuentes consultadas", level=1)
    for fuente in pack.fuentes:
        doc.add_paragraph(fuente, style="List Bullet")

    doc.save(destino)


# ── 2 · evaluación de riesgos ────────────────────────────────────────────────────────────


def _hoja(libro: Workbook, titulo: str, columnas: list[tuple[str, int]]) -> None:
    hoja = libro.create_sheet(titulo) if libro.sheetnames != ["Sheet"] else libro.active
    hoja.title = titulo

    for indice, (nombre, ancho) in enumerate(columnas, start=1):
        celda = hoja.cell(row=1, column=indice, value=nombre)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = CABECERA
        celda.alignment = Alignment(vertical="center", wrap_text=True)
        hoja.column_dimensions[get_column_letter(indice)].width = ancho

    hoja.freeze_panes = "A2"
    return hoja


def evaluacion_riesgos(destino: Path, pack) -> None:
    """
    La plantilla del sistema de gestión de riesgos del art. 9.

    Va en Excel porque es lo que se revisa cada trimestre y se compara con la revisión anterior;
    en Word nadie compara dos versiones de una tabla.
    """
    libro = Workbook()

    hoja = _hoja(libro, "Riesgos", [
        ("Id", 8), ("Sistema", 24), ("Riesgo identificado", 40),
        ("¿A quién afecta?", 22), ("Probabilidad (1-5)", 14), ("Impacto (1-5)", 12),
        ("Nivel", 10), ("Medida de mitigación", 40), ("Riesgo residual", 14),
        ("Responsable", 20), ("Revisión", 14),
    ])

    ejemplos = pack.riesgos_ejemplo
    for fila, (riesgo, afectado, medida) in enumerate(ejemplos, start=2):
        hoja.cell(row=fila, column=1, value=f"R-{fila - 1:02d}")
        hoja.cell(row=fila, column=3, value=riesgo)
        hoja.cell(row=fila, column=4, value=afectado)
        hoja.cell(row=fila, column=8, value=medida)

        # El nivel se calcula, no se teclea: dos personas rellenando a mano acaban puntuando
        # distinto el mismo riesgo.
        hoja.cell(row=fila, column=7, value=f"=IF(E{fila}*F{fila}>=15,\"Alto\","
                                            f"IF(E{fila}*F{fila}>=6,\"Medio\",\"Bajo\"))")

    guia = libro.create_sheet("Cómo se usa")
    guia["A1"] = "Sistema de gestión de riesgos"
    guia["A1"].font = Font(bold=True, size=14)

    lineas = [
        "",
        f"Base: art. 9 del {fuentes.AI_ACT}. El sistema de gestión de riesgos es un proceso",
        "continuo y planificado a lo largo de todo el ciclo de vida, no un documento que se firma",
        "una vez.",
        "",
        "Las filas de ejemplo son de este sector y están para borrarlas: sirven para ver el nivel",
        "de concreción que hace falta. «El modelo puede fallar» no es un riesgo identificado.",
        "",
        f"Verificación normativa: {fuentes.FECHA_VERIFICACION}",
        f"Versión: {fuentes.VERSION}",
        "",
        fuentes.AVISO,
    ]
    for indice, texto in enumerate(lineas, start=2):
        guia.cell(row=indice, column=1, value=texto)

    guia.column_dimensions["A"].width = 100

    libro.save(destino)


# ── 3 · registro de logs ─────────────────────────────────────────────────────────────────


def registro_logs(destino: Path, pack) -> None:
    """
    Qué hay que registrar y cuánto hay que guardarlo.

    La parte que ahorra dinero de verdad: casi todo esto ya lo produce un sistema con trazas y
    evaluaciones. El registro no es infraestructura nueva, es ponerle nombre y plazo a lo que ya
    se emite.
    """
    libro = Workbook()

    hoja = _hoja(libro, "Registro", [
        ("Acontecimiento", 34), ("Base normativa", 18), ("¿De dónde sale hoy?", 30),
        ("Campos mínimos", 44), ("Conservación", 16), ("Responsable", 20), ("¿Implantado?", 14),
    ])

    filas = [
        ("Cada uso del sistema", "art. 12.1",
         "Traza de la petición",
         "Identificador, inicio, fin, versión del sistema y del modelo", "≥ 6 meses"),
        ("Datos de entrada de la decisión", "art. 12.2.c / 26.4",
         "Traza de la petición",
         "Entradas, o su referencia si son datos personales que no deban duplicarse", "≥ 6 meses"),
        ("Resultado y su confianza", "art. 12.2.a",
         "Traza de la petición",
         "Salida, puntuación, umbral aplicado", "≥ 6 meses"),
        ("Intervención humana", "art. 14 / 26.2",
         "Registro de la revisión",
         "Quién revisó, cuándo, si confirmó o corrigió, y el motivo", "≥ 6 meses"),
        ("Anomalía o deriva detectada", "art. 12.2.a / 72",
         "Evaluaciones periódicas",
         "Métrica, valor, umbral, ventana temporal", "≥ 6 meses"),
        ("Incidente grave", "art. 26.5 / 73",
         "Gestión de incidentes",
         "Qué pasó, cuándo se detectó, a quién se informó y cuándo", "≥ 6 meses"),
        ("Cambio de versión del sistema o del modelo", "art. 12.2.a",
         "Despliegue",
         "Versión anterior, nueva, fecha, responsable", "≥ 6 meses"),
    ]

    for indice, (evento, base, origen, campos, plazo) in enumerate(filas, start=2):
        hoja.cell(row=indice, column=1, value=evento)
        hoja.cell(row=indice, column=2, value=base)
        hoja.cell(row=indice, column=3, value=origen)
        hoja.cell(row=indice, column=4, value=campos)
        hoja.cell(row=indice, column=5, value=plazo)
        hoja.cell(row=indice, column=7, value="☐")

        for columna in range(1, 8):
            hoja.cell(row=indice, column=columna).alignment = Alignment(wrap_text=True, vertical="top")

    notas = libro.create_sheet("Plazos y matices")
    notas["A1"] = "Lo que hay que saber antes de fijar el plazo"
    notas["A1"].font = Font(bold=True, size=14)

    lineas = ["", "Obligaciones de quien despliega el sistema:"]
    lineas += [f"  {ref} · {texto}" for ref, texto in fuentes.OBLIGACIONES_DESPLIEGUE
               if ref in ("26.5", "26.6", "26.11")]
    lineas += ["", "Qué tiene que permitir registrar el sistema:"]
    lineas += [f"  {ref} · {texto}" for ref, texto in fuentes.REGISTROS_ART12]
    lineas += ["", *pack.notas_registro, "",
               f"Verificación normativa: {fuentes.FECHA_VERIFICACION}", "", fuentes.AVISO]

    for indice, texto in enumerate(lineas, start=2):
        celda = notas.cell(row=indice, column=1, value=texto)
        celda.alignment = Alignment(wrap_text=True, vertical="top")

    notas.column_dimensions["A"].width = 110

    libro.save(destino)


# ── 4 · cláusulas modelo ─────────────────────────────────────────────────────────────────


def clausulas_modelo(destino: Path, pack) -> None:
    """
    El reparto de responsabilidades con el proveedor del modelo, por escrito.

    Lo que más se ignora: se puede pasar a ser PROVEEDOR sin pretenderlo, por poner la marca
    propia o por cambiar la finalidad. Y entonces se heredan las obligaciones del art. 16.
    """
    doc = Document()
    _portada_docx(doc, "Cláusulas modelo proveedor · responsable del despliegue", pack.sector)

    doc.add_paragraph(
        "Redactado desde la posición de quien DESPLIEGA el sistema y contrata a un proveedor. "
        "Los corchetes son lo que hay que rellenar; lo demás es articulado."
    )

    doc.add_heading("Por qué hace falta un acuerdo escrito", level=1)
    doc.add_paragraph(fuentes.ACUERDO_ESCRITO_25_4)

    doc.add_heading("Cuándo el que despliega pasa a ser proveedor", level=1)
    doc.add_paragraph(
        "Las tres circunstancias del art. 25.1. Si se da alguna, se asumen las obligaciones del "
        "art. 16 y el proveedor inicial deja de serlo para ese sistema:"
    )
    for referencia, texto in fuentes.CONVERTIRSE_EN_PROVEEDOR:
        _cita(doc, referencia, texto)

    doc.add_paragraph(
        "La letra a) admite pacto en contrario —«sin perjuicio de los acuerdos contractuales que "
        "estipulen que las obligaciones se asignan de otro modo»—. Las letras b) y c) no: ahí no "
        "hay cláusula que valga, porque quien modifica sustancialmente el sistema o le cambia la "
        "finalidad es quien sabe lo que ha hecho."
    )

    doc.add_heading("Cláusulas", level=1)

    clausulas = [
        ("Objeto y finalidad prevista",
         "El PROVEEDOR declara que la finalidad prevista del sistema es [FINALIDAD]. EL CLIENTE "
         "se obliga a no destinarlo a una finalidad distinta sin acuerdo escrito previo, dado que "
         "un cambio de finalidad puede convertirle en proveedor conforme al art. 25.1.c."),
        ("Clasificación y su documentación",
         "El PROVEEDOR entregará su evaluación de clasificación conforme al art. 6 y, si concluye "
         "que el sistema no es de alto riesgo pese a estar en el Anexo III, la documentación "
         "exigida por el art. 6.4."),
        ("Instrucciones de uso",
         "El PROVEEDOR entregará las instrucciones de uso con la información del art. 13, "
         "actualizadas ante cualquier cambio, en [IDIOMA] y con [PLAZO] de antelación."),
        ("Información para la supervisión humana",
         "El PROVEEDOR facilitará lo necesario para que EL CLIENTE cumpla el art. 26.2, "
         "incluyendo las limitaciones conocidas del sistema y las señales por las que un "
         "supervisor puede detectar que la salida no es fiable."),
        ("Registros",
         "El sistema generará los archivos de registro del art. 12 y EL CLIENTE podrá acceder a "
         "ellos y exportarlos en [FORMATO] durante toda la vigencia y [PLAZO] después. La "
         "conservación mínima aplicable es de seis meses (art. 26.6), sin perjuicio de plazos "
         "mayores exigidos por otra normativa."),
        ("Incidentes",
         "El PROVEEDOR notificará a EL CLIENTE cualquier incidente grave o riesgo del art. 79.1 "
         "en un plazo máximo de [PLAZO] desde que tenga conocimiento, con la información "
         "necesaria para que EL CLIENTE cumpla el art. 26.5."),
        ("Cambios sustanciales",
         "El PROVEEDOR informará con [PLAZO] de antelación de cualquier modificación sustancial "
         "del sistema o del modelo subyacente. EL CLIENTE podrá resolver sin penalización si la "
         "modificación altera la clasificación de riesgo o las condiciones de supervisión."),
        ("Cooperación con autoridades",
         "El PROVEEDOR cooperará con las autoridades competentes y facilitará a EL CLIENTE la "
         "documentación que estas le requieran, en el plazo que estas fijen."),
        ("Datos de entrenamiento y gobernanza de datos",
         "El PROVEEDOR acreditará el cumplimiento del art. 10 respecto de los conjuntos de datos "
         "empleados, en la medida en que resulte aplicable al sistema contratado."),
        ("Subcontratación",
         "El PROVEEDOR informará de los terceros que intervengan en la prestación y de su "
         "ubicación, y notificará por adelantado cualquier cambio."),
    ]

    for indice, (titulo, texto) in enumerate(clausulas, start=1):
        doc.add_heading(f"{indice}. {titulo}", level=2)
        doc.add_paragraph(texto)

    if pack.clausulas_sector:
        doc.add_heading("Cláusulas adicionales de este sector", level=1)
        doc.add_paragraph(pack.clausulas_sector_intro)

        for indice, (titulo, texto) in enumerate(pack.clausulas_sector, start=len(clausulas) + 1):
            doc.add_heading(f"{indice}. {titulo}", level=2)
            doc.add_paragraph(texto)

    doc.add_heading("Fuentes consultadas", level=1)
    for fuente in pack.fuentes:
        doc.add_paragraph(fuente, style="List Bullet")

    doc.save(destino)


# ── 5 · caso resuelto ────────────────────────────────────────────────────────────────────


def caso_resuelto(destino: Path, pack) -> None:
    """
    El pack aplicado a un caso concreto, de principio a fin.

    Es el documento que hace que los otros cuatro se entiendan: enseña una clasificación real
    con su conclusión y, sobre todo, POR QUÉ. Va en PDF porque se lee, no se rellena.
    """
    estilos = getSampleStyleSheet()
    cuerpo = ParagraphStyle(
        "cuerpo", parent=estilos["BodyText"], alignment=TA_JUSTIFY, spaceAfter=8, leading=14)
    titulo = ParagraphStyle(
        "titulo", parent=estilos["Title"], fontSize=18, spaceAfter=6, alignment=0)
    seccion = ParagraphStyle(
        "seccion", parent=estilos["Heading2"], fontSize=13, spaceBefore=14, spaceAfter=6)
    menudo = ParagraphStyle("menudo", parent=estilos["BodyText"], fontSize=8, textColor="#555555")

    partes = [
        Paragraph("INKOOVA ACADEMY", menudo),
        Paragraph(pack.caso_titulo, titulo),
        Paragraph(
            f"Pack sectorial · {pack.sector} · versión {fuentes.VERSION}<br/>"
            f"<b>Verificación normativa: {fuentes.FECHA_VERIFICACION}</b>", cuerpo),
        Paragraph(fuentes.AVISO, menudo),
        Spacer(1, 0.6 * cm),
    ]

    for encabezado, parrafos in pack.caso:
        partes.append(Paragraph(encabezado, seccion))

        for parrafo in parrafos:
            if isinstance(parrafo, list):
                partes.append(ListFlowable(
                    [ListItem(Paragraph(p, cuerpo)) for p in parrafo],
                    bulletType="bullet", start="•", leftIndent=14))
            else:
                partes.append(Paragraph(parrafo, cuerpo))

    partes.append(PageBreak())
    partes.append(Paragraph("Fuentes consultadas", seccion))
    partes.append(ListFlowable(
        [ListItem(Paragraph(f, cuerpo)) for f in pack.fuentes],
        bulletType="bullet", start="•", leftIndent=14))

    SimpleDocTemplate(
        str(destino), pagesize=A4,
        leftMargin=2.2 * cm, rightMargin=2.2 * cm, topMargin=2 * cm, bottomMargin=2 * cm,
        title=pack.caso_titulo, author="Inkoova Academy",
    ).build(partes)
